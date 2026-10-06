"""
Amazon Flex MTR → SAP B1 Consolidated Posting
=============================================
Amazon Flex is not integrated with Unicommerce, so B2C revenue from the
monthly MTR consolidation file is posted to SAP B1 as consolidated
AR Invoices (and AR Credit Notes for refunds), mirroring the
Unicommerce → SAP B1 B2C accounting treatment.

Configuration (same as Amazon_IN_API channel):
  Customer Code  : C03575
  Series         : 500 invoices / 507 credit notes (FY 2026-27)
  SGST Account   : 41106001  (intra-state: Karnataka → Karnataka)
  IGST Account   : 41106002  (inter-state: Karnataka → other states)
  Warehouse      : EC-FG
  U_BillingFrom  : KT (always Karnataka — ships from PRCJ Bengaluru)

Design:
  - B2C and B2B rows are merged; split is by GST bucket only:
      ONE consolidated AR Invoice for intra-state (CGST+SGST) shipments
      ONE consolidated AR Invoice for inter-state (IGST) shipments
      ONE consolidated AR Credit Note per tax bucket for refunds
    → 2 invoices + up to 2 credit notes per file
  - GST filing happens from the MTR itself
  - DocDate is supplied by the user (not taken from the file)
  - Lines are consolidated per (FG Code, Tax Code); UnitPrice = total
    taxable value / total quantity so document totals match the file
  - CFG parent items are expanded via SAP ProductTrees (Sales BOM) and
    batch numbers are auto-allocated FEFO from EC-FG stock
  - FreeReplacement rows are skipped

Usage — preview totals and payloads (no posting):
    bench --site dev.localhost execute \\
        khanal_tech_integrations.utils.MTR.amazon_flex_conso_sap.preview \\
        --args '("/path/to/KFPL_Amazon Conso file_June 2026 (2).xlsx", "2026-06-30")'

Usage — post the two consolidated AR Invoices:
    bench --site dev.localhost execute \\
        khanal_tech_integrations.utils.MTR.amazon_flex_conso_sap.post_invoices \\
        --args '("/path/to/file.xlsx", "2026-06-30")'

Usage — post the consolidated AR Credit Notes (refund rows):
    bench --site dev.localhost execute \\
        khanal_tech_integrations.utils.MTR.amazon_flex_conso_sap.post_credit_notes \\
        --args '("/path/to/file.xlsx", "2026-06-30")'

Usage — post everything (invoices first, then credit notes):
    bench --site dev.localhost execute \\
        khanal_tech_integrations.utils.MTR.amazon_flex_conso_sap.post_all \\
        --args '("/path/to/file.xlsx", "2026-06-30")'
"""

import json
import logging

import frappe
import openpyxl

from khanal_tech_integrations.utils.sap import AuthenticateSAPB1

logger = logging.getLogger(__name__)

HEADERS = {
    "Accept": "*/*",
    "User-Agent": "Khanal Tech",
    "Content-Type": "application/json",
}

CONFIG = {
    "customer_code":      "C03575",
    "account_code_sgst":  "41106001",
    "account_code_igst":  "41106002",
    "series_invoice":     "500",
    "series_credit_note": "507",
    "warehouse_code":     "EC-FG",
    "billing_from_state": "KT",
}

BILL_TO_CODE = {
    "IGST": "B2C IGST ADD",
    "SGST": "B2C SGST ADD",
}

# GST-rate → SAP tax code, per bucket
TAX_CODE_MAP = {
    "SGST": {0: "KACS0", 5: "KACS5", 12: "KACS12", 18: "KACS18"},
    "IGST": {0: "KAIG0", 5: "KAIG5", 12: "KAIG12", 18: "KAIG18"},
}

SHEET_NAME = "Flex"


# ---------------------------------------------------------------------------
# SAP session helpers
# ---------------------------------------------------------------------------

_session_cache = {"session": None}
_product_tree_cache = {}
_sap_item_cache = {}
_sap_batch_stock_cache = {}


def _get_session():
    if not _session_cache["session"]:
        _session_cache["session"] = AuthenticateSAPB1()
    return _session_cache["session"]


def _get_sap_url():
    sap_b1_url = frappe.get_doc("SAP Settings").sap_b1_url or ""
    return sap_b1_url.rstrip("/") + "/"


def _clean(value):
    return str(value or "").strip()


def _fnum(value):
    try:
        return float(value or 0)
    except (TypeError, ValueError):
        return 0.0


def _is_cfg_parent_item(item_code):
    return _clean(item_code).upper().startswith("CFG")


def _sap_yes(value):
    return str(value or "").strip().upper() in {"TYES", "Y", "YES", "TRUE", "1"}


def _get_sap_item_flags(item_code):
    item_code = _clean(item_code)
    if not item_code:
        return {}
    if item_code in _sap_item_cache:
        return _sap_item_cache[item_code]

    session = _get_session()
    b1_url = _get_sap_url()
    item_esc = item_code.replace("'", "''")
    endpoints = [
        f"Items('{item_esc}')?$select=ItemCode,InventoryItem,ManageBatchNumbers,ManageSerialNumbers",
        f"Items?$select=ItemCode,InventoryItem,ManageBatchNumbers,ManageSerialNumbers&$filter=ItemCode eq '{item_esc}'",
    ]

    for endpoint in endpoints:
        try:
            resp = session.request("GET", f"{b1_url}{endpoint}", headers=HEADERS, verify=False)
            if resp.status_code != 200:
                continue

            data = resp.json() or {}
            if "value" in data:
                items = data.get("value") or []
                data = items[0] if items else {}

            if data.get("ItemCode"):
                if not any(field in data for field in ("InventoryItem", "ManageBatchNumbers", "ManageSerialNumbers")):
                    continue
                flags = {
                    "inventory_item": _sap_yes(data.get("InventoryItem")),
                    "batch_managed": _sap_yes(data.get("ManageBatchNumbers")),
                    "serial_managed": _sap_yes(data.get("ManageSerialNumbers")),
                }
                _sap_item_cache[item_code] = flags
                return flags
        except Exception as e:
            logger.warning(f"Unable to fetch SAP item flags for {item_code}: {e}")

    _sap_item_cache[item_code] = {}
    return {}


def _requires_batch_numbers(item_code):
    flags = _get_sap_item_flags(item_code)
    if flags.get("serial_managed"):
        raise ValueError(f"Item {item_code} is serial-managed; only batch items are supported")
    if flags:
        return flags.get("inventory_item") and flags.get("batch_managed")
    raise ValueError(f"Unable to determine SAP batch-management flags for item {item_code}")


def _get_sap_batch_stock_rows(item_code, include_zero=False):
    from khanal_tech_integrations.api.common.sap_connector import SAPConnector

    warehouse_code = CONFIG["warehouse_code"]
    item_code = _clean(item_code)
    cache_key = (warehouse_code, item_code, include_zero)
    if cache_key in _sap_batch_stock_cache:
        return _sap_batch_stock_cache[cache_key]

    qty_filter = "" if include_zero else 'q."Quantity" > 0 AND'
    connector = SAPConnector()
    connection = cursor = None
    try:
        connection, cursor, schema = connector.get_hana_connection()
        cursor.execute(
            f"""
            SELECT b."ItemCode", b."DistNumber", q."WhsCode", q."Quantity", b."ExpDate"
            FROM {schema}."OBTN" b
            JOIN {schema}."OBTQ" q
                ON b."ItemCode" = q."ItemCode" AND b."SysNumber" = q."SysNumber"
            WHERE {qty_filter} q."WhsCode" = ? AND b."ItemCode" = ?
            ORDER BY b."ExpDate", b."DistNumber"
            """,
            [warehouse_code, item_code],
        )
        sap_rows = cursor.fetchall() or []
    finally:
        connector.close_hana_connection(connection, cursor)

    rows = [
        {
            "distnumber": _clean(row[1]),
            "quantity": _fnum(row[3]),
            "expdate": str(row[4]) if row[4] else "9999-12-31",
        }
        for row in sap_rows
    ]
    _sap_batch_stock_cache[cache_key] = rows
    return rows


def _allocate_batch_numbers(item_code, quantity, reservations, enforce_stock=True, invoice_batches=None):
    """FEFO batch allocation from EC-FG. For credit notes (incoming stock,
    enforce_stock=False) the quantity is received back into the same
    batches the posted invoices shipped from (invoice_batches); items not
    on those invoices fall back to the newest batch the warehouse held."""
    item_code = _clean(item_code)
    required_qty = _fnum(quantity)
    if required_qty <= 0:
        return []
    if not _requires_batch_numbers(item_code):
        return []

    if not enforce_stock:
        if invoice_batches:
            allocation = []
            remaining = required_qty
            for batch_no, invoiced_qty in invoice_batches:
                if remaining <= 0:
                    break
                take = min(remaining, invoiced_qty)
                allocation.append({"BatchNumber": batch_no, "Quantity": round(take, 6)})
                remaining = round(remaining - take, 6)
            if remaining > 0:
                # refund exceeds what those batches shipped; rest joins the last batch
                allocation[-1]["Quantity"] = round(allocation[-1]["Quantity"] + remaining, 6)
            return allocation
        # item not on the posted invoices (e.g. refund of an earlier month):
        # any batch the warehouse has ever held is valid for receiving
        batch_rows = _get_sap_batch_stock_rows(item_code, include_zero=True)
        if not batch_rows:
            raise ValueError(f"No batch found in {CONFIG['warehouse_code']} for credit-note item {item_code}")
        batch = sorted(batch_rows, key=lambda r: r["expdate"])[-1]
        return [{"BatchNumber": batch["distnumber"], "Quantity": round(required_qty, 6)}]

    batch_rows = _get_sap_batch_stock_rows(item_code)
    remaining = required_qty
    batch_numbers = []
    for row in batch_rows:
        batch_no = row["distnumber"]
        if not batch_no:
            continue
        reservation_key = (item_code, batch_no)
        available = row["quantity"] - reservations.get(reservation_key, 0.0)
        if available <= 0:
            continue
        assigned = min(available, remaining)
        batch_numbers.append({"BatchNumber": batch_no, "Quantity": round(assigned, 6)})
        reservations[reservation_key] = reservations.get(reservation_key, 0.0) + assigned
        remaining = round(remaining - assigned, 6)
        if remaining <= 0:
            break

    if remaining > 0:
        raise ValueError(
            f"Insufficient batch stock for {item_code} in {CONFIG['warehouse_code']}: "
            f"required {required_qty}, short by {remaining}"
        )
    return batch_numbers


def _get_product_tree_children(parent_item_code):
    parent_item_code = _clean(parent_item_code)
    if parent_item_code in _product_tree_cache:
        return _product_tree_cache[parent_item_code]

    session = _get_session()
    b1_url = _get_sap_url()
    item_esc = parent_item_code.replace("'", "''")
    last_error = None

    for endpoint in (
        f"ProductTrees('{item_esc}')",
        f"ProductTrees('{item_esc}')?$expand=ProductTreeLines",
        f"ProductTrees?$filter=TreeCode eq '{item_esc}'&$expand=ProductTreeLines",
    ):
        try:
            resp = session.request("GET", f"{b1_url}{endpoint}", headers=HEADERS, verify=False)
            if resp.status_code != 200:
                last_error = f"SAP {resp.status_code}: {resp.text[:200]}"
                continue
            data = resp.json() or {}
            if "value" in data:
                trees = data.get("value") or []
                data = trees[0] if trees else {}
            lines = data.get("ProductTreeLines") or []
            if isinstance(lines, dict):
                lines = lines.get("value") or []
            if not lines:
                # older Service Layers return neither inline nor $expand lines
                lines_resp = session.request(
                    "GET",
                    f"{b1_url}ProductTrees('{item_esc}')/ProductTreeLines",
                    headers=HEADERS,
                    verify=False,
                )
                if lines_resp.status_code == 200:
                    lines_data = lines_resp.json() or {}
                    lines = lines_data.get("value", lines_data if isinstance(lines_data, list) else [])
            children = [
                {"ItemCode": _clean(l.get("ItemCode")), "Quantity": _fnum(l.get("Quantity"))}
                for l in lines
                if _clean(l.get("ItemCode")) and _fnum(l.get("Quantity"))
            ]
            if children:
                _product_tree_cache[parent_item_code] = children
                return children
            last_error = f"ProductTree {parent_item_code} has no child lines"
        except Exception as e:
            last_error = str(e)

    raise ValueError(
        f"Unable to resolve Sales BOM children for CFG parent {parent_item_code}. "
        f"Last error: {last_error}"
    )


# ---------------------------------------------------------------------------
# MTR file parsing
# ---------------------------------------------------------------------------

def _parse_flex_sheet(file_path):
    """Read the Flex sheet, return list of dicts (cached formula values)."""
    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    if SHEET_NAME not in wb.sheetnames:
        raise ValueError(f"Sheet '{SHEET_NAME}' not found in {file_path}")
    ws = wb[SHEET_NAME]

    header = None
    rows = []
    for raw in ws.iter_rows(values_only=True):
        if header is None:
            if raw and raw[0] == "Type":
                header = list(raw)
            continue
        if raw[0] is None:
            continue
        rows.append(dict(zip(header, raw)))
    if header is None:
        raise ValueError("Header row (first cell 'Type') not found in Flex sheet")
    return rows


def _bucket_of(row):
    """intra (SGST) vs inter (IGST) from Amazon's own tax columns."""
    if _fnum(row.get("Cgst Tax")) or _fnum(row.get("Sgst Tax")):
        return "SGST"
    if _fnum(row.get("Igst Tax")):
        return "IGST"
    # zero-tax rows: fall back on destination state
    return "SGST" if _clean(row.get("Ship To State")).upper() == "KARNATAKA" else "IGST"


def _gst_rate_pct(row):
    rate = _fnum(row.get("GST Rate"))
    if rate == 0:
        rate = (
            _fnum(row.get("Cgst Rate")) + _fnum(row.get("Sgst Rate")) + _fnum(row.get("Igst Rate"))
        )
    return int(round(rate * 100))


def _taxable_value(row):
    taxable = row.get("Taxable Value")
    if taxable is None:
        taxable = row.get("Tax Exclusive Gross")
    return _fnum(taxable)


# Amazon has used different literal values for the same real-world event
# across export vintages ("Shipment" in older MTR files, "Sales" seen from
# July 2026) — both mean a shipped sale that should invoice. Refund has
# only ever had the one value so far, kept as an exact match.
_TXN_TYPE_SYNONYMS = {
    "Shipment": {"SHIPMENT", "SALE", "SALES"},
    "Refund": {"REFUND"},
}


def _matches_txn_type(row_txn_type, txn_type):
    normalized = _clean(row_txn_type).upper()
    synonyms = _TXN_TYPE_SYNONYMS.get(txn_type)
    if synonyms:
        return normalized in synonyms
    return normalized == txn_type.upper()


# FG codes that are broken catalog duplicates of a real bundle, same root
# cause as Unicommerce Settings.sku_alias_map on the ops-center posting
# side (khanal_tech_integrations/unicommerce/mappers/sap_payload.py).
# FGHN0056 ("Pack of 3") isn't SAP-registered as a Sales BOM (TreeType='P',
# no ProductTree lines), so the CFG-parent expansion below never applies
# to it — it needs a literal item+qty substitution instead. Taxable value
# is left untouched; unit_price is derived downstream as taxable/qty, so
# multiplying qty here automatically divides the value evenly across the
# real component units.
FG_CODE_ALIASES = {
    "FGHN0056": {"item_code": "FGHN0054", "qty_multiplier": 3},
}


def _consolidate(rows, txn_type, per_order=False):
    """
    Group B2C + B2B rows of one transaction type into the two tax
    buckets. Returns {"SGST": {key: {...}}, "IGST": {...}} with summed
    qty / taxable. Default key is (FG Code, tax code) — fully
    consolidated; with per_order=True the key includes the Amazon invoice
    number so each order stays its own document line (U_Order traceable
    in SAP). Every group also records its source Amazon invoice numbers.
    Refund rows carry negative amounts but positive quantities in the
    file; both are normalised to positive for the credit note.
    """
    is_refund = txn_type == "Refund"
    buckets = {"SGST": {}, "IGST": {}}

    for row in rows:
        if _clean(row.get("Type")) not in ("B2C", "B2B"):
            continue
        if not _matches_txn_type(row.get("Transaction Type"), txn_type):
            continue

        fg_code = _clean(row.get("FG Code"))
        if not fg_code:
            raise ValueError(f"Missing FG Code on invoice {row.get('Invoice Number')}")

        alias = FG_CODE_ALIASES.get(fg_code)
        qty_multiplier = alias["qty_multiplier"] if alias else 1
        if alias:
            fg_code = alias["item_code"]

        bucket = _bucket_of(row)
        rate = _gst_rate_pct(row)
        tax_code = TAX_CODE_MAP[bucket].get(rate)
        if not tax_code:
            raise ValueError(f"No SAP tax code for GST rate {rate}% ({bucket}) on {row.get('Invoice Number')}")

        qty = (abs(_fnum(row.get("Quantity"))) if is_refund else _fnum(row.get("Quantity"))) * qty_multiplier
        taxable = abs(_taxable_value(row)) if is_refund else _taxable_value(row)
        order_no = _clean(row.get("Credit Note No") if is_refund else row.get("Invoice Number")) \
            or _clean(row.get("Invoice Number"))

        key = (order_no, fg_code, tax_code) if per_order else (fg_code, tax_code)
        entry = buckets[bucket].setdefault(
            key, {"fg_code": fg_code, "tax_code": tax_code, "qty": 0.0, "taxable": 0.0, "orders": set()}
        )
        entry["qty"] += qty
        entry["taxable"] += taxable
        entry["orders"].add(order_no)

    # drop groups that net out to nothing
    for bucket in buckets:
        buckets[bucket] = {
            k: v for k, v in buckets[bucket].items() if round(v["qty"], 6) > 0 and round(v["taxable"], 2) > 0
        }
    return buckets


# ---------------------------------------------------------------------------
# Payload builders
# ---------------------------------------------------------------------------

def _build_sap_line(line_num, item_code, quantity, unit_price, tax_code, account_code,
                    doc_label, tree_type=None, batch_reservations=None, enforce_stock=True,
                    cn_batch_map=None):
    line = {
        "LineNum":       line_num,
        "ItemCode":      item_code,
        "Quantity":      round(quantity, 6),
        "UnitPrice":     unit_price,
        "AccountCode":   account_code,
        "WarehouseCode": CONFIG["warehouse_code"],
        "TaxCode":       tax_code,
        "TaxType":       "tt_Yes",
        "TaxLiable":     "tYES",
        "U_Order":       doc_label,
        "U_Country":     "India",
    }
    if tree_type:
        line["TreeType"] = tree_type
    if tree_type != "iSalesTree" and batch_reservations is not None:
        invoice_batches = (cn_batch_map or {}).get(_clean(item_code))
        batch_numbers = _allocate_batch_numbers(item_code, quantity, batch_reservations, enforce_stock, invoice_batches)
        if batch_numbers:
            line["BatchNumbers"] = batch_numbers
    return line


def _expand_group_to_lines(group, account_code, start_line_num, doc_label,
                           batch_reservations, enforce_stock=True, cn_batch_map=None):
    """One consolidated (FG, tax code) group → SAP lines (CFG → BOM children)."""
    item_code = group["fg_code"]
    quantity = group["qty"]
    unit_price = round(group["taxable"] / quantity, 6) if quantity else 0.0
    tax_code = group["tax_code"]
    orders = sorted(group.get("orders") or [])
    if len(orders) == 1:
        doc_label = orders[0]

    if not _is_cfg_parent_item(item_code):
        return [
            _build_sap_line(start_line_num, item_code, quantity, unit_price, tax_code,
                            account_code, doc_label,
                            batch_reservations=batch_reservations, enforce_stock=enforce_stock,
                            cn_batch_map=cn_batch_map)
        ], start_line_num + 1

    # Sales BOMs here price by COMPONENTS: SAP zeroes the iSalesTree
    # parent's price, so the group's whole value must ride on the child
    # lines (verified June 2026 — parent-priced docs lost the value).
    # Split proportionally by quantity across ALL children at a uniform
    # per-unit price, not dumped on the first child alone (that older
    # version priced every child after the first at 0.0, which fails
    # SAP's "Unit Price cannot be less than Item Cost" check outright
    # for any bundle with more than one distinct component — see the
    # 2026-08 Flex July run). The last child absorbs the rounding
    # remainder so total line value still matches the group's taxable
    # value exactly, same convention as the simple-item unit_price calc
    # above.
    lines = [
        _build_sap_line(start_line_num, item_code, quantity, 0.0, tax_code,
                        account_code, doc_label, tree_type="iSalesTree",
                        batch_reservations=batch_reservations, enforce_stock=enforce_stock,
                        cn_batch_map=cn_batch_map)
    ]
    line_num = start_line_num + 1
    children = _get_product_tree_children(item_code)
    child_qtys = [round(quantity * child["Quantity"], 6) for child in children]
    total_units = sum(child_qtys)
    per_unit_price = round(group["taxable"] / total_units, 6) if total_units else 0.0

    remaining_value = group["taxable"]
    for idx, (child, child_qty) in enumerate(zip(children, child_qtys)):
        is_last = idx == len(children) - 1
        if is_last:
            child_price = round(remaining_value / child_qty, 6) if child_qty else 0.0
        else:
            child_price = per_unit_price
            remaining_value = round(remaining_value - child_price * child_qty, 6)
        lines.append(
            _build_sap_line(line_num, child["ItemCode"], child_qty, child_price,
                            tax_code, account_code, doc_label, tree_type="iIngredient",
                            batch_reservations=batch_reservations, enforce_stock=enforce_stock,
                            cn_batch_map=cn_batch_map)
        )
        line_num += 1
    return lines, line_num


def _num_at_card(bucket, is_credit_note, doc_date):
    """Unique customer-ref per document; used as the idempotency key."""
    return f"AMZFLEX-{bucket}{'-CN' if is_credit_note else ''}-{doc_date}"


def _build_doc_payload(bucket, groups, doc_date, doc_label, is_credit_note, batch_reservations, cn_batch_map=None):
    series = CONFIG["series_credit_note"] if is_credit_note else CONFIG["series_invoice"]
    bill_to_code = BILL_TO_CODE[bucket]
    account_code = CONFIG["account_code_sgst"] if bucket == "SGST" else CONFIG["account_code_igst"]

    payload = {
        "CardCode":                    CONFIG["customer_code"],
        "NumAtCard":                   _num_at_card(bucket, is_credit_note, doc_date),
        "Comments":                    (
            f"Amazon Flex B2C+B2B {'Intra-state' if bucket == 'SGST' else 'Inter-state'} "
            f"consolidated {'Credit Note' if is_credit_note else 'Invoice'} — {doc_label}"
        ),
        "PayToCode":                   bill_to_code,
        "ShipToCode":                  bill_to_code,
        "DocDate":                     doc_date,
        "DocDueDate":                  doc_date,
        "U_BillingFrom":               CONFIG["billing_from_state"],
        "U_BillTo":                    "Local" if bucket == "SGST" else "Central",
        "TransportationCode":          5,
        "UseBillToAddrToDetermineTax": "tNO",
        "Series":                      series,
        "DocumentLines":               [],
    }

    line_num = 0
    for key in sorted(groups):
        lines, line_num = _expand_group_to_lines(
            groups[key], account_code, line_num, doc_label,
            batch_reservations, enforce_stock=not is_credit_note,
            cn_batch_map=cn_batch_map,
        )
        payload["DocumentLines"].extend(lines)
    return payload


# ---------------------------------------------------------------------------
# Posting
# ---------------------------------------------------------------------------

def _get_invoice_batch_map(doc_date):
    """
    item → [(batch_no, qty), ...] from the two posted June AR Invoices,
    so credit notes receive stock back into the same batches shipped.
    """
    session = _get_session()
    b1_url = _get_sap_url()
    batch_map = {}
    for bucket in ("SGST", "IGST"):
        existing = _find_existing_sap_doc("Invoices", _num_at_card(bucket, False, doc_date))
        if not existing:
            continue
        resp = session.request(
            "GET", f"{b1_url}Invoices({existing['DocEntry']})", headers=HEADERS, verify=False
        )
        resp.raise_for_status()
        for line in (resp.json() or {}).get("DocumentLines") or []:
            item = _clean(line.get("ItemCode"))
            for b in line.get("BatchNumbers") or []:
                batch_no = _clean(b.get("BatchNumber"))
                qty = _fnum(b.get("Quantity"))
                if item and batch_no and qty:
                    batch_map.setdefault(item, []).append((batch_no, qty))
    return batch_map


def _find_existing_sap_doc(endpoint, num_at_card):
    """Return {DocEntry, DocNum} if a doc with this NumAtCard already exists (idempotency guard)."""
    session = _get_session()
    url = (
        f"{_get_sap_url()}{endpoint}"
        f"?$select=DocEntry,DocNum&$filter=CardCode eq '{CONFIG['customer_code']}' "
        f"and NumAtCard eq '{num_at_card}' and Cancelled eq 'tNO'"
    )
    try:
        resp = session.request("GET", url, headers=HEADERS, verify=False)
        resp.raise_for_status()
        docs = (resp.json() or {}).get("value") or []
        return docs[0] if docs else None
    except Exception as e:
        raise RuntimeError(f"Duplicate check failed for {num_at_card}: {e}") from e


def _save_results(results, endpoint, doc_date):
    """Persist run results as JSON under the site's private files; returns path."""
    import os
    from datetime import datetime

    try:
        out_dir = frappe.get_site_path("private", "files", "amazon_flex_mtr")
        os.makedirs(out_dir, exist_ok=True)
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        path = os.path.join(out_dir, f"{endpoint.lower()}_{doc_date}_{stamp}.json")
        with open(path, "w") as fh:
            json.dump(results, fh, indent=2, default=str)
        print(f"[RESULTS SAVED] {path}")
        return path
    except Exception as e:
        logger.warning(f"Could not save results file: {e}")
        return None


def _post_to_sap(endpoint, payload, doc_label):
    session = _get_session()
    url = f"{_get_sap_url()}{endpoint}"
    try:
        resp = session.request("POST", url, data=json.dumps(payload), headers=HEADERS, verify=False)
        resp.raise_for_status()
        result = resp.json()
        msg = f"[SAP OK] {doc_label} → DocNum {result.get('DocNum')}, DocEntry {result.get('DocEntry')}"
        logger.info(msg)
        print(msg)
        return result
    except Exception as e:
        error_body = ""
        if getattr(e, "response", None) is not None:
            error_body = e.response.text
        msg = f"[SAP FAIL] {doc_label} → {e} | {error_body}"
        logger.error(msg)
        print(msg)
        return None


def _bucket_totals(groups):
    all_orders = set()
    for g in groups.values():
        all_orders |= g.get("orders") or set()
    return {
        "line_groups": len(groups),
        "total_qty": round(sum(g["qty"] for g in groups.values()), 2),
        "total_taxable": round(sum(g["taxable"] for g in groups.values()), 2),
        "source_orders": len(all_orders),
    }


def _order_map(groups):
    """FG+tax code → sorted Amazon invoice numbers, for the results JSON audit trail."""
    return {
        f"{g['fg_code']}|{g['tax_code']}": sorted(g.get("orders") or [])
        for g in groups.values()
    }


def _run(file_path, doc_date, txn_type, endpoint, dry_run, per_order=False):
    if not doc_date or len(str(doc_date)) != 10:
        raise ValueError("doc_date must be 'YYYY-MM-DD' (user-selected posting date)")

    rows = _parse_flex_sheet(file_path)
    is_cn = endpoint == "CreditNotes"

    results = {}
    buckets = _consolidate(rows, txn_type, per_order=per_order)
    doc_label = f"AMZ FLEX MTR {'CN ' if is_cn else ''}{doc_date}"
    # one reservation ledger across ALL documents in this run, so the
    # SGST and IGST invoices can never double-book the same batch stock
    batch_reservations = {}
    # credit notes receive stock back into the batches the posted
    # invoices shipped from
    cn_batch_map = _get_invoice_batch_map(str(doc_date)) if is_cn else None

    for bucket in ("SGST", "IGST"):
        groups = buckets[bucket]
        summary = _bucket_totals(groups)
        if not groups:
            results[bucket] = {"status": "skipped (no rows)", **summary}
            continue

        ref = _num_at_card(bucket, is_cn, str(doc_date))
        existing = None if dry_run else _find_existing_sap_doc(endpoint, ref)
        if existing:
            results[bucket] = {
                "status": "already posted (skipped)",
                **summary,
                "NumAtCard": ref,
                "DocEntry": existing.get("DocEntry"),
                "DocNum": existing.get("DocNum"),
            }
            print(f"[SKIP] {bucket} already in SAP as DocNum {existing.get('DocNum')} ({ref})")
            continue

        payload = _build_doc_payload(bucket, groups, str(doc_date), doc_label, is_cn, batch_reservations, cn_batch_map)
        if dry_run:
            results[bucket] = {"status": "preview", **summary, "NumAtCard": ref, "payload": payload}
            continue

        label = f"{'CN' if is_cn else 'INV'} {bucket} {doc_label}"
        sap_result = _post_to_sap(endpoint, payload, label)
        results[bucket] = {
            "status": "posted" if sap_result else "FAILED",
            **summary,
            "NumAtCard": ref,
            "DocEntry": (sap_result or {}).get("DocEntry"),
            "DocNum": (sap_result or {}).get("DocNum"),
            "order_map": _order_map(groups),
        }

    if not dry_run:
        _save_results(results, endpoint, str(doc_date))
    return results


def check_stock(file_path, doc_date):
    """
    Report EC-FG batch-stock shortages for the invoices still pending.
    Sums required qty per item across the unposted invoice buckets
    (CFG parents expanded to BOM children) and compares against total
    available batch stock. Prints one line per short item. Posts nothing.
    """
    rows = _parse_flex_sheet(file_path)
    buckets = _consolidate(rows, "Shipment")

    need = {}
    for bucket in ("SGST", "IGST"):
        groups = buckets[bucket]
        if not groups:
            continue
        ref = _num_at_card(bucket, False, str(doc_date))
        existing = _find_existing_sap_doc("Invoices", ref)
        if existing:
            print(f"[SKIP] {bucket} invoice already posted (DocNum {existing.get('DocNum')}) — not counted")
            continue
        for g in groups.values():
            item = g["fg_code"]
            qty = g["qty"]
            if _is_cfg_parent_item(item):
                for child in _get_product_tree_children(item):
                    need[child["ItemCode"]] = need.get(child["ItemCode"], 0.0) + qty * child["Quantity"]
            else:
                need[item] = need.get(item, 0.0) + qty

    shortages = []
    for item in sorted(need):
        required = need[item]
        if not _requires_batch_numbers(item):
            continue
        available = sum(r["quantity"] for r in _get_sap_batch_stock_rows(item))
        if available < required:
            shortages.append({
                "item": item,
                "required": round(required, 2),
                "available_ec_fg": round(available, 2),
                "short_by": round(required - available, 2),
            })

    if not shortages:
        print("NO SHORTAGES — all pending invoice items covered by EC-FG batch stock")
    else:
        print(f"{len(shortages)} item(s) short in EC-FG:")
        for s in shortages:
            print(f"  {s['item']}: need {s['required']}, have {s['available_ec_fg']}, SHORT {s['short_by']}")
    return shortages


def preview(file_path, doc_date, per_order=0):
    """Dry run: parse, consolidate, build payloads, print totals. Nothing posted."""
    per_order = bool(int(per_order))
    out = {
        "invoices": _run(file_path, doc_date, "Shipment", "Invoices", dry_run=True, per_order=per_order),
        "credit_notes": _run(file_path, doc_date, "Refund", "CreditNotes", dry_run=True, per_order=per_order),
    }
    for kind, buckets in out.items():
        for bucket, info in buckets.items():
            print(kind, bucket, {k: v for k, v in info.items() if k != "payload"})
    return out


def post_invoices(file_path, doc_date, per_order=0):
    """Post the two consolidated AR Invoices (B2C+B2B merged): intra (SGST+CGST) and inter (IGST).

    per_order=1 keeps one document line per Amazon invoice number
    (U_Order = PRCJ-XXXX) instead of merging lines per FG code."""
    results = _run(file_path, doc_date, "Shipment", "Invoices", dry_run=False, per_order=bool(int(per_order)))
    print(json.dumps(results, indent=2, default=str))
    return results


def post_credit_notes(file_path, doc_date, per_order=0):
    """Post the consolidated AR Credit Notes (B2C+B2B merged) for Refund rows."""
    results = _run(file_path, doc_date, "Refund", "CreditNotes", dry_run=False, per_order=bool(int(per_order)))
    print(json.dumps(results, indent=2, default=str))
    return results


def post_all(file_path, doc_date, per_order=0):
    """Invoices first, then credit notes."""
    return {
        "invoices": post_invoices(file_path, doc_date, per_order),
        "credit_notes": post_credit_notes(file_path, doc_date, per_order),
    }
